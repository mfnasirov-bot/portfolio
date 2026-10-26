"""AI Invoice Extractor: reads PDF invoices with Claude and writes the data to Excel.

Usage:
    python invoice_extractor.py invoices/ -o invoices.xlsx
"""

import argparse
import base64
import sys
from pathlib import Path
from typing import Optional

import anthropic
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from pydantic import BaseModel, Field

MODEL = "claude-opus-5-5"

PROMPT = """Extract the key data from this invoice.

Rules:
- Dates must be in YYYY-MM-DD format.
- total_amount is the final amount due (including tax), as a number without currency symbols.
- currency is a 3-letter ISO code (e.g. USD, EUR, AZN).
- If a field is missing, unreadable, or ambiguous, leave it null, set needs_review to true,
  and explain why in review_reason. Never guess a number."""


class Invoice(BaseModel):
    vendor_name: Optional[str] = Field(description="Company that issued the invoice")
    invoice_number: Optional[str]
    invoice_date: Optional[str] = Field(description="YYYY-MM-DD")
    due_date: Optional[str] = Field(description="YYYY-MM-DD")
    total_amount: Optional[float]
    currency: Optional[str]
    needs_review: bool
    review_reason: Optional[str]


COLUMNS = [
    ("File", "file"),
    ("Vendor", "vendor_name"),
    ("Invoice #", "invoice_number"),
    ("Invoice date", "invoice_date"),
    ("Due date", "due_date"),
    ("Total", "total_amount"),
    ("Currency", "currency"),
    ("Needs review", "needs_review"),
    ("Review reason", "review_reason"),
]

REVIEW_FILL = PatternFill("solid", start_color="FFF2CC")


def extract_invoice(client: anthropic.Anthropic, pdf_path: Path) -> Invoice:
    pdf_data = base64.standard_b64encode(pdf_path.read_bytes()).decode("ascii")
    response = client.beta.messages.parse(
        model=MODEL,
        max_tokens=4096,
        output_config={"effort": "low"},
        betas=["server-side-fallback-2026-07-01"],
        fallbacks="default",
        messages=[{
            "role": "user",
            "content": [
                {
                    "type": "document",
                    "source": {"type": "base64", "media_type": "application/pdf", "data": pdf_data},
                },
                {"type": "text", "text": PROMPT},
            ],
        }],
        output_format=Invoice,
    )
    if response.stop_reason == "refusal" or response.parsed_output is None:
        return Invoice(
            vendor_name=None, invoice_number=None, invoice_date=None, due_date=None,
            total_amount=None, currency=None, needs_review=True,
            review_reason=f"Could not extract (stop reason: {response.stop_reason})",
        )
    return response.parsed_output


def write_excel(rows: list[dict], output_path: Path) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Invoices"
    ws.append([title for title, _ in COLUMNS])
    for cell in ws[1]:
        cell.font = Font(bold=True)

    for row in rows:
        ws.append([
            ("YES" if row[key] else "") if key == "needs_review" else row[key]
            for _, key in COLUMNS
        ])
        if row["needs_review"]:
            for cell in ws[ws.max_row]:
                cell.fill = REVIEW_FILL

    for column in ws.columns:
        width = max(len(str(cell.value or "")) for cell in column)
        ws.column_dimensions[column[0].column_letter].width = min(width + 2, 60)
    ws.freeze_panes = "A2"
    wb.save(output_path)


def main() -> int:
    parser = argparse.ArgumentParser(description="Extract invoice data from PDFs into Excel.")
    parser.add_argument("input", type=Path, help="A PDF file or a folder of PDFs")
    parser.add_argument("-o", "--output", type=Path, default=Path("invoices.xlsx"))
    args = parser.parse_args()

    pdfs = sorted(args.input.glob("*.pdf")) if args.input.is_dir() else [args.input]
    if not pdfs:
        print(f"No PDF files found in {args.input}")
        return 1

    client = anthropic.Anthropic()
    rows = []
    for i, pdf in enumerate(pdfs, 1):
        print(f"[{i}/{len(pdfs)}] {pdf.name} ... ", end="", flush=True)
        try:
            invoice = extract_invoice(client, pdf)
        except anthropic.AuthenticationError:
            print("\nInvalid API key. Set ANTHROPIC_API_KEY and try again.")
            return 1
        except TypeError as e:
            if "authentication method" not in str(e):
                raise
            print("\nNo API key found. Set ANTHROPIC_API_KEY first (see README.md).")
            return 1
        except anthropic.APIConnectionError:
            print("\nNetwork error. Check your internet connection.")
            return 1
        except anthropic.APIStatusError as e:
            print(f"API error ({e.status_code}): {e.message}")
            invoice = Invoice(
                vendor_name=None, invoice_number=None, invoice_date=None, due_date=None,
                total_amount=None, currency=None, needs_review=True,
                review_reason=f"API error {e.status_code}",
            )
        else:
            status = "needs review" if invoice.needs_review else "ok"
            print(f"{invoice.vendor_name} | {invoice.total_amount} {invoice.currency} | {status}")
        rows.append({"file": pdf.name, **invoice.model_dump()})

    write_excel(rows, args.output)
    flagged = sum(r["needs_review"] for r in rows)
    print(f"\nDone: {len(rows)} invoices -> {args.output} ({flagged} flagged for review)")
    return 0


if __name__ == "__main__":
    sys.exit(main())
